# modules/pages/10_state_helpers.py -- one segment of the CivilProposals app script.
# Session-state defaults and every shared helper: job-queue-or-inline runner, project identity/key, paid-status checks, structure rebuild, save/load/autosave.
#
# HOW THIS FILE RUNS: app.py executes the files in modules/pages/ in order,
# in ONE shared script namespace (see app.py's loader) -- exactly as when
# all of this was a single 315KB app.py, just split along its natural seams
# so each area of the app lives in its own reviewable file. Code here is
# UNCHANGED from the pre-split app.py; names defined in earlier segments
# (helpers, current_user, _access, the tabs list) are used directly, and
# st.session_state remains the single shared state, same as before.
# Deliberately NOT a function-per-page refactor: that would change scoping
# and evaluation order (the thing a mid-tender regression hides in), and is
# better done tab-by-tab with live click-through testing.
from __future__ import annotations
# ---------------------------------------------------------------------------
# Session state helpers
# ---------------------------------------------------------------------------

def _init_state():
    defaults = {
        "project_name": "", "client_name": "", "tender_name": "",
        "submission_date_input": "", "bidder_name": "", "proposal_theme": "Corporate",
        "project_type": PROJECT_TYPES[0],
        # Language of the AI-GENERATED proposal content itself (drafts, executive
        # summary, team intro, etc.) -- a per-PROJECT choice, independent of the
        # app's own UI language (modules/i18n.py's "_lang"). See Project Setup tab.
        "output_language": "en",
        "tender_extracted": None,
        "company_material_text": {},
        "company_material_files": {},  # {category: {filename: extracted_text}} -- per-file, so a
                                        # re-upload can update/remove individual files (see the
                                        # Upload Documents tab) instead of replacing the whole category.
        "company_uploaded_flags": {},
        "returnable_schedule_files": {},  # {filename: raw bytes} -- returnable schedules from a package ZIP (see package_intake)
        "project_photo_bytes": [],
        # Which uploaded photo goes on the cover. Was hardcoded to the first
        # one, with no way to choose -- the most visible image in the pack
        # decided by upload order.
        "cover_photo_index": 0,
        "branding_bytes": [],
        # Pre-filled from ANTHROPIC_API_KEY in a local .env file, if present (see
        # load_dotenv() at the top of this file) -- so this tab doesn't need to be
        # re-filled in on every launch. Still fully editable/overridable below.
        "ai_config": {
            "provider": "Anthropic Claude",
            "api_key": _ENV_ANTHROPIC_KEY,
            "model": ai_interface.get_default_model("Anthropic Claude") if _ENV_ANTHROPIC_KEY else "",
            "endpoint": "",
        },
        "copilot_client_id": "", "copilot_tenant_id": "",
        "copilot_access_token": "", "copilot_token_cache": "", "copilot_username": "",
        # Whether the sidebar's Claude API key field should persist to the local
        # .env file -- defaults to on if a key was already loaded from .env at
        # startup (i.e. it was remembered before), off otherwise.
        "_remember_claude_key": bool(_ENV_ANTHROPIC_KEY),
        "analysis": None,
        "weighted_criteria": None,
        "allocations": None,
        "sections": None,
        "guidance_notes": None,
        "compliance_items": None,
        "gap_items": None,
        "drafts": None,
        "executive_summary": None,
        "team_intro": None,
        "experience_intro": None,
        "project_differentiator": "",
        "project_sales_pitch": "",
        "pitch_review": None,
        "pitch_questions": None,
        "tender_summary_buffer": None,
        "graphics": None,
        "weighting_chart_png": None,
        "fee_estimates": None,
        "docx_buffer": None,
        "quotes": [],
        "section_divider_config": {},
        "divider_images": {},
        "cover_hero_png": None,
        # Small Scope pack (short, brief-driven response pack) -- one more option within
        # the same pipeline, not a separate app. See PROPOSAL_FORMAT_LABELS above.
        "proposal_format": "formal",
        # Which proposal_format st.session_state.sections was actually built under --
        # see _structure_format_stale(). None until Proposal Structure (tab 4) is
        # generated at least once.
        "_sections_built_format": None,
        # Sign-off details only -- the cover page/footer already carry project, client,
        # and bidder details, so no separate recipient/ref/date/subject fields are needed.
        "letter_sender_name": "", "letter_sender_title": "",
        "letter_sender_phone": "", "letter_sender_email": "",
        # Read by returnable_schedules.build_fill_data for the "Registered
        # Office"/"Business Address" labels on the client's own forms.
        # Deliberately NOT part of the letter sign-off block -- see the note
        # in export_docx._build_letter_signoff.
        "letter_sender_address": "",
        "terms_of_engagement_text": "",
        "team_members": [],
        "team_photos": {},
        "team_bio_warnings": [],
        # Resourcing plan + org chart (Team & Resourcing tab). resource_plan is a
        # list of resourcing.ResourceAssignment; org_chart_png is the generated PNG.
        "resource_plan": [],
        "resource_extra_names": [],
        "dismissed_disciplines": [],
        # AI-suggested "include in proposal" verdicts for the Key Personnel pen pics
        # (Team & Resourcing tab) -- {slot: {"recommended": bool, "reason": str}},
        # from resourcing.suggest_proposal_inclusion(). Empty until the "Suggest
        # which personnel to include" button is run; purely advisory -- the actual
        # include/exclude state lives on each ResourceAssignment.include_in_proposal.
        "personnel_inclusion_suggestions": {},
        "cv_library_filenames": [],
        "cv_extracted_names": [],
        "org_chart_png": None,
        # Which of the four presentation styles the org chart is drawn in
        # (org_chart_render.STYLES), and which style the stored PNG above was
        # actually generated in. The second key is what lets the tab say "the
        # exported pack still has the previous style" instead of silently
        # exporting something other than what is on screen.
        "org_chart_style": org_chart_render.DEFAULT_STYLE,
        "org_chart_png_style": "",
        # Optional management roles (resourcing.OPTIONAL_MANAGEMENT_ROLES) the
        # user has removed from this project's chart. Kept as an explicit
        # record rather than inferred from the plan's absence, for the same
        # reason dismissed_disciplines exists: without it, every reconcile
        # pass would helpfully put the role back.
        "removed_management_roles": [],
        "body_font": "Arial",
        # Key-personnel headshots (Team & Resourcing tab), keyed by person_name --
        # same pattern as team_photos. None means "not run yet" is not applicable
        # here (it's always a dict, possibly empty).
        "personnel_photos": {},
        # Reference projects (Upload Documents tab) -- structured Section 2 content
        # drafted/revised from the "Project references" upload; see
        # modules/reference_projects.py. None means the drafting step hasn't run
        # yet (distinct from "ran and found nothing"); reference_project_photos is
        # always a dict, keyed by reference project title.
        "reference_projects": None,
        "reference_project_photos": {},
        "reference_project_warnings": [],
        # First-pass manual discipline fee build-up (Fee tab). List of
        # resourcing.DisciplineFeeLine; always includes Project Management.
        "discipline_fee_lines": [],
        # Disciplines the user explicitly removed from the Fee Estimate table --
        # never re-added by the brief-sync merge (separate from
        # dismissed_disciplines, which is the Team & Resourcing tab's own list).
        "dismissed_fee_disciplines": [],
        # Bump counters folded into the discipline/scope-item fee data_editor
        # widget keys below -- st.data_editor with a fixed key ignores its
        # `data` argument on every rerun after the first (the widget owns its
        # state once created), so merging a newly-extracted discipline or
        # scope item into the underlying session_state list alone is not
        # enough: the editor keeps showing whatever it first rendered,
        # forever, even across a fresh Tender Analysis run. Bumping the
        # version whenever a genuinely new row gets merged in forces a fresh
        # widget instance (which picks up the merged data) without touching
        # the key on every ordinary rerun, which would instead throw away
        # in-progress edits. See the merge blocks on the Fee Estimate tab.
        "_discipline_fee_editor_version": 0,
        "_scope_fee_editor_version": 0,
        "_large_scope_fee_editor_version": 0,
        # Cache for the discipline fee tables' Excel export + pie chart, so a
        # fragment rerun that doesn't change the underlying hours/rates
        # doesn't waste time (or widen the edit-commit race window) redoing
        # that work -- see the discipline fee table fragments below.
        "_disc_fee_cache_sig": None,
        "_disc_fee_cache_xlsx": None,
        "_disc_fee_cache_pie": None,
        "_letter_disc_fee_cache_sig": None,
        "_letter_disc_fee_cache_xlsx": None,
        "_letter_disc_fee_cache_pie": None,
        # Deferred-apply state for the fee tables -- the rebuild (dedup/dismiss
        # logic) and the cache above only run when the user presses this
        # table's "Apply changes" button (see _fee_apply_control), rather than
        # on every keystroke-commit. Only the last-applied signature is kept
        # now: the tick/tick_seen pair existed solely to make a checkbox
        # behave like a button, which a button does on its own.
        "_disc_fee_last_applied_editor_sig": None,
        "_letter_disc_fee_last_applied_editor_sig": None,
        # Same deferred-apply pattern, extended to the other three fee-editing
        # tables (scope item / deliverable fee build-up, both pack sizes, and
        # the discipline fee % split, both pack sizes).
        "_scope_fee_last_applied_editor_sig": None,
        "_large_scope_fee_last_applied_editor_sig": None,
        "_pct_fee_last_applied_editor_sig": None,
        "_letter_pct_fee_last_applied_editor_sig": None,
        "scope_item_fees": [],
        # Which of the three fee presentations go into the proposal body.
        # The app builds up to three in parallel -- the indicative % split by
        # discipline, the hours x rate discipline build-up, and the scope-item
        # build-up -- and which of them appeared was hardcoded per pack format
        # with no way to choose. These defaults reproduce exactly what both
        # packs exported before the choice existed: % split + discipline
        # build-up in, scope-item build-up out (it was never exported by
        # either pack, and its own on-screen caption calls it internal
        # tracking).
        "fee_sections_included": {
            "pct_split": True,
            "discipline_buildup": True,
            "scope_buildup": False,
        },
        "fee_seed_total": 0.0,
        # Manually-entered total project fee for the indicative benchmark split
        # below -- overrides the brief's stated fee cap (if any) so the user can
        # still see a $ split even when the brief never states a ceiling.
        "fee_estimate_manual_total": 0.0,
        # Small Scope pack's own total for its "Discipline fee split (%)" table --
        # separate from fee_estimate_manual_total (Large Scope's own, unrelated
        # override) since this one auto-prepopulates from the discipline fee
        # build-up total the first time it's used (0.0 = "not yet set"), then
        # stays independently editable. See the Fee Estimate tab.
        "letter_fee_total_override": 0.0,
        # The reviewed methodology stage grid (methodology_stages.MethodologyStage
        # list). None means "the stage drafter has never been run" -- distinct
        # from [] which would mean "run, and produced nothing".
        "methodology_stages": None,
        # Which of the four methodology presentation styles (see
        # methodology_render.STYLES) drives the exported PowerPoint and its
        # live preview -- same selector pattern as program_style and
        # org_chart_style below.
        "methodology_style": methodology_render.DEFAULT_STYLE,
        # First-pass risk/impact/mitigation table (risk_register.RiskRegister).
        # None means the step has never been run.
        "risk_register": None,
        # Whether the user has confirmed their firm issues Work Verification
        # Records. Never assumed: the WVR line used to print as fact in every
        # methodology export without the app ever being told.
        "methodology_wvr_confirmed": False,
        "program_num_weeks": 6,
        "program_schedule": {},
        "program_week_labels": [],
        # The user's own anticipated start date for the delivery program
        # (datetime.date | None). Optional: set it and every week header
        # becomes a real calendar date via program_schedule.week_labels().
        # Never derived from the brief -- a guessed start date in a program
        # table would be an invented fact.
        "program_start_date": None,
        # Which of the four presentation styles the delivery program is drawn
        # in (program_render.STYLES). A presentation preference, not
        # brief-derived data -- so it is deliberately NOT reset when the
        # brief changes, the same reasoning as fee_sections_included.
        "program_style": program_render.DEFAULT_STYLE,
        # Save/Load Project bookkeeping (sidebar) -- not project content itself.
        "_project_save_bytes": None,
        "_last_loaded_project_name": "",
        "_autosave_enabled": True,
        "_last_autosave_ts": 0.0,
        "_last_autosave_path": "",
        "_last_autosave_error": "",
    }
    for key, value in defaults.items():
        if key not in st.session_state:
            st.session_state[key] = value


def _run_job_or_inline(job_type, func, args=(), kwargs=None, progress=None,
                        queued_text="Queued...", running_text="Working...",
                        inline_extra_kwargs=None, queue_func=None, queue_args=None):
    """
    Runs func(*args, **kwargs) either in the background job queue (see
    modules/job_queue.py) or inline in this process, and either way blocks
    until it's done and returns func's return value -- every call site
    keeps its existing "click button, wait, see result" behaviour. The only
    thing that changes is WHERE the AI-call/CPU time is actually spent.

    The queue is used only when all of these are true: this is a logged-in
    SaaS user (current_user is set -- local/dev use with SAAS_MODE=false
    has no account to own a job), AND a Redis-backed queue is actually
    configured (job_queue.redis_available()) -- i.e. the worker service
    from DEPLOY.md's "Background jobs" section has been deployed. Otherwise
    this falls back to calling func directly, exactly as every one of these
    call sites always worked before background jobs existed, using
    inline_extra_kwargs for anything that only makes sense in-process (a
    live progress_callback closure can't be pickled across the process
    boundary to a worker, so it's never passed to the queued path).

    queue_func/queue_args let a call site enqueue a DIFFERENT function/args
    than the ones it runs inline with -- used by both current call sites to
    enqueue a job_queue.run_*_job() wrapper with a redacted ai_config
    (api_key="") instead of func/args' real one, so the real server
    Anthropic key never ends up pickled into Redis (see job_queue.py's
    docstring). Falls back to func/args when not given, so this stays a
    no-op for any future call site that doesn't need the distinction.

    Deliberately never a hard dependency on the queue: turning it on is
    just setting REDIS_URL and deploying the worker service, nothing here
    has to change, and this app never breaks just because that hasn't
    happened yet.
    """
    kwargs = kwargs or {}
    use_queue = IS_SAAS_MODE and current_user and job_queue.redis_available()

    if not use_queue:
        return func(*args, **kwargs, **(inline_extra_kwargs or {}))

    enqueue_func = queue_func or func
    enqueue_args = args if queue_args is None else queue_args
    enqueue_kwargs = dict(kwargs)
    if queue_func is not None:
        # Both job_queue.run_*_job wrappers accept usage_context -- plain
        # strings attributing the job's AI calls to this user/project for
        # per-bid cost logging (db.AiCallLog). Only added when a wrapper is
        # in use: a future call site that enqueues its function directly
        # shouldn't get a kwarg it never asked for.
        enqueue_kwargs["usage_context"] = {
            "user_id": current_user.id,
            "project_key": _current_project_key(),
            "project_name": st.session_state.get("project_name", ""),
            "purpose": job_type,
        }
    job_id = job_queue.enqueue(current_user.id, job_type, enqueue_func, *enqueue_args, **enqueue_kwargs)
    if progress:
        progress.progress(0.05, text=queued_text)

    # Hard ceiling on how long this waits, in case a job is enqueued but
    # never picked up (e.g. the worker service -- see modules/worker.py --
    # is down or was never deployed) or somehow never resolves. Without
    # this, that "queued" status is indistinguishable from "still working"
    # forever, and the caller's `while True` would poll indefinitely with
    # the progress bar stuck creeping toward 90% -- which, to whoever's
    # sitting there watching their one free trial bid, looks exactly like
    # the product being broken, not like a transient infra problem. Set
    # comfortably above job_queue.JOB_TIMEOUT_SECONDS (RQ's own 15-minute
    # execution ceiling once a worker actually starts the job) so a
    # legitimately long-running job is never preempted before RQ itself
    # would have given up on it.
    MAX_WAIT_SECONDS = 20 * 60
    elapsed = 0.0
    poll_interval = 1.5
    while elapsed < MAX_WAIT_SECONDS:
        status = job_queue.get_status(job_id, current_user.id)
        if status["status"] == "finished":
            return status["result"]
        if status["status"] in ("failed", "not_found"):
            raise RuntimeError(status["error"] or "The background job failed.")
        if progress:
            # Indeterminate-ish: creeps toward 90% over ~3 minutes rather
            # than claiming a precision this polling loop doesn't have --
            # queued jobs have no live per-chunk/per-section count to show
            # (see this function's docstring), just "still working".
            progress.progress(min(0.05 + elapsed / 180.0, 0.9), text=running_text)
        time.sleep(poll_interval)
        elapsed += poll_interval

    raise RuntimeError(
        "This is taking much longer than expected and may mean the background worker is "
        "offline. Nothing was charged for this attempt beyond what the trial/subscription "
        "gate already checked -- try again in a few minutes, and contact "
        "hello@civilproposals.com if it keeps happening."
    )


def _ocr_export_note() -> str | None:
    """The "OCR-derived -- verify carefully" notice for exported documents
    (see export_docx._add_ocr_notice), or None when none of the brief's text
    came from OCR. Kept here so all four export call sites share one wording."""
    _ext = st.session_state.get("tender_extracted")
    if not getattr(_ext, "ocr_used", False):
        return None
    pages = getattr(_ext, "ocr_pages", []) or []
    page_part = f" (page {', '.join(str(p) for p in pages[:12])}{' and more' if len(pages) > 12 else ''})" if pages else ""
    return (
        f"Parts of the tender brief this pack is based on were read with text recognition "
        f"(OCR) from scanned pages{page_part}. {document_processor.OCR_VERIFY_TAG}: OCR can "
        f"misread numbers, dates, and names -- verify extracted requirements, figures, and "
        f"names against the original brief before this document goes anywhere near a "
        f"submission. Delete this notice once verified."
    )


def _project_info() -> dict:
    return {
        "project_name": st.session_state.project_name,
        "client_name": st.session_state.client_name,
        "tender_name": st.session_state.tender_name,
        "submission_date": st.session_state.submission_date_input,
        "bidder_name": st.session_state.bidder_name,
        "proposal_theme": st.session_state.proposal_theme,
        "project_type": st.session_state.project_type,
    }


def _current_project_key() -> str:
    """Same identity/hash scheme as the Tender Analysis tab's own
    _project_key (see that tab's comment for why the brief's own text is
    hashed into it, not just the typed names) -- pulled out here so every
    downstream/auxiliary AI feature can check the SAME project's paid
    status without duplicating this logic at every call site."""
    _brief_text = st.session_state.tender_extracted.text if st.session_state.tender_extracted else ""
    _brief_hash = hashlib.sha256(_brief_text.encode("utf-8")).hexdigest()[:16] if _brief_text else ""
    return (
        f"{st.session_state.project_name}|{st.session_state.tender_name}|"
        f"{st.session_state.client_name}|{_brief_hash}"
    ).strip("|")


def _current_project_already_paid() -> bool:
    """True once THIS project has actually had a bid spent on it via
    auth.record_proposal_usage() -- i.e. Tender Analysis has genuinely run
    (and been paid/trialled/subscribed-for) for it, not just that the
    account currently HOLDS some unspent capacity somewhere.

    Every downstream/auxiliary AI feature below (CV name extraction,
    discipline re-scan, team bios, fee benchmark refresh, draft generation,
    exec summary, pitch review) gates on THIS, in addition to or instead of
    _access["allowed"], for two opposite reasons that both trace back to
    the same fix:
      - _access["allowed"] alone was too PERMISSIVE: it only means "this
        account has some capacity somewhere" and never changes just because
        one of these particular calls is made (nothing here decrements
        anything the way Tender Analysis does) -- so an account sitting on
        one never-spent trial bid could hammer CV extraction / discipline
        re-scan / bios / benchmark refresh an unlimited number of times,
        forever, without ever actually running (or paying for) a Tender
        Analysis on any project.
      - _access["allowed"] alone was also too RESTRICTIVE right after the
        trial bid actually gets spent: the instant Tender Analysis consumes
        the account's one trial bid, _access["allowed"] flips to False for
        the whole account, which used to block draft generation / exec
        summary / pitch review even on the SAME project that bid was just
        spent analysing -- so a trial user could see an analysis and then
        immediately hit a paywall before ever seeing a draft, contradicting
        the "your first full bid gets you a drafted proposal pack" promise
        (see the welcome email).
    Checking "has THIS specific project already been paid for" instead of
    "does the account have spare capacity right now" fixes both at once:
    once you've paid to analyse a project, everything else about producing
    THAT proposal is included, and nothing about a project you haven't paid
    for yet is.

    Unlimited accounts (auth.UNLIMITED_ACCOUNTS) and non-SaaS/local use
    always pass -- same as everywhere else that checks _access.

    Also folds in the account-level AI-spend ceiling and fair-use rate
    limit (see modules/limits.py, and _ai_gate_msg computed once per
    script run in 00_init.py) -- both are checked here rather than added
    as a separate condition at each of the dozen-plus call sites, so a
    project that WAS paid for can still correctly go back to "AI not
    available right now" once the account's trial spend ceiling is hit,
    without touching every button's `ready` boolean individually. Use
    _ai_block_reason() below for the matching help/caption text -- it
    gives the specific reason (spend ceiling / rate limit / not-yet-paid)
    this function collapses into a single bool."""
    if not IS_SAAS_MODE or not current_user:
        return True
    if _access.get("unlimited"):
        return True
    if _ai_gate_msg:
        return False
    _key = _current_project_key()
    if not _key:
        return False
    with db.get_session() as s:
        return s.query(db.ProposalUsage).filter(
            db.ProposalUsage.user_id == current_user.id,
            db.ProposalUsage.project_key == _key.lower(),
        ).first() is not None


def _ai_block_reason() -> str | None:
    """Matching help/caption text for _current_project_already_paid()'s
    account-level checks (see that function's docstring) -- highest
    priority first: the AI-spend ceiling or fair-use rate limit message
    (account-wide, applies no matter what project this is), else
    _PROJECT_NOT_PAID_HINT once neither of those is the reason. Returns
    None when _current_project_already_paid() is actually True (nothing to
    explain). Every AI-feature help/caption site that used to write
    `_PROJECT_NOT_PAID_HINT if not _current_project_already_paid() else
    "<feature-specific hint>"` should use `_ai_block_reason() if not
    _current_project_already_paid() else "<feature-specific hint>"`
    instead, so the ceiling/rate-limit message actually gets seen instead
    of a stale 'run Tender Analysis first' hint that no longer applies."""
    if _ai_gate_msg:
        return _ai_gate_msg
    return _PROJECT_NOT_PAID_HINT


# ---------------------------------------------------------------------------
# Part B (one-pass free tier) / Part B2 (pass allowances per plan) --
# entitlement checks for Export Pack's artifact downloads and Tender
# Analysis's generation passes. See auth.project_funded_by(),
# auth.is_trial_funded_project(), auth.project_passes_status() and
# db.ArtifactEvent / db.ProjectPasses for the underlying data model.
# ---------------------------------------------------------------------------

# The exactly-3 artifacts a free-trial project gets ONE download each of
# (Part B). Everything else Export Pack can produce (Methodology Table
# PPTX, Delivery Program PPTX, filled returnable schedules) is paid-only,
# full stop -- never offered to a trial-funded project regardless of
# download count. Keys match the `artifact_type` values written to
# db.ArtifactEvent by _mark_free_artifact_downloaded() below and read by
# modules/pages/80_export.py's download buttons.
FREE_TIER_ARTIFACTS = ("proposal_docx", "tender_summary_docx", "org_chart_pptx")


def _project_is_free_tier() -> bool:
    """True when the CURRENT project is subject to Part B's free-tier
    artifact rules (a trial-funded, or not-yet-funded, project) -- False
    for a paid (subscription/credit) or UNLIMITED_ACCOUNTS project, and
    always False outside SaaS mode (the desktop/BYOK build has no
    billing concept at all, so nothing here ever restricts it)."""
    if not IS_SAAS_MODE or not current_user:
        return False
    if _access.get("unlimited"):
        return False
    _key = _current_project_key()
    if not _key:
        # No project identity yet -- nothing has been (or can be) analysed
        # or exported, so there's nothing to gate; the readiness checklist
        # elsewhere already blocks getting this far without one.
        return False
    return auth.is_trial_funded_project(current_user, _key)


_TRIAL_ARTIFACT_EVENT_SENTINEL_KEY = ""


def _free_artifact_already_downloaded(artifact_type: str) -> bool:
    """True once this ACCOUNT's ONE free download of `artifact_type` (one
    of FREE_TIER_ARTIFACTS) has already happened. Only meaningful while
    _project_is_free_tier() is True -- a paid project never writes or
    checks these rows at all (see db.ArtifactEvent's docstring).

    Audit fix Part 2b: checked by (user_id, artifact_type) alone, NOT also
    by project_key -- previously scoping by project_key meant the free
    download was really "once per project IDENTITY", and project identity
    folds in a hash of the brief text (see _current_project_key()), so
    editing even one character of the project/tender/client name computed
    a brand new, never-downloaded-from identity with no ArtifactEvent rows
    against it: a fresh "one free download" on demand, repeatable forever
    on the trial tier. The trial is fundamentally ONE bid; "one free
    download of each artifact on the trial, ever" (account-wide) is the
    correct rule and closes the rename bypass outright. Deliberately
    ignores project_key on the read side too, so an ArtifactEvent row
    written under the OLD per-project scheme (before this fix shipped)
    still correctly counts as "already used" -- no data migration needed,
    see _mark_free_artifact_downloaded()'s sentinel-key approach below."""
    if not _project_is_free_tier():
        return False
    with db.get_session() as s:
        return s.query(db.ArtifactEvent).filter(
            db.ArtifactEvent.user_id == current_user.id,
            db.ArtifactEvent.artifact_type == artifact_type,
        ).first() is not None


def _mark_free_artifact_downloaded(artifact_type: str) -> None:
    """Records that this ACCOUNT's one free download of `artifact_type`
    has now happened. Call this the moment a free-tier download actually
    fires (see modules/pages/80_export.py's on_click= callbacks) -- never
    for a paid project (see _project_is_free_tier()), which has no
    download cap to track.

    Audit fix Part 2b: writes with project_key=_TRIAL_ARTIFACT_EVENT_SENTINEL_KEY
    (a fixed empty string, not this project's real key) so the EXISTING
    unique constraint on (user_id, project_key, artifact_type) enforces
    "once per account", not "once per project" -- see
    _free_artifact_already_downloaded()'s docstring for why per-project
    scoping was a rename-driven bypass, and why this needs no database
    schema change: a real project's key is always non-empty (a project
    name is required before Tender Analysis can even run), so the sentinel
    can never collide with one.

    Audit fix Part 2c: previously caught a bare `except Exception` and
    silently swallowed it -- meaning ANY database hiccup on this insert
    (not just the expected duplicate-key "already downloaded" case) failed
    OPEN, granting an unmetered extra download with no record and no sign
    anything went wrong. Now only a genuine duplicate-key IntegrityError
    (the harmless, expected "already downloaded" outcome -- including a
    concurrent click racing this exact insert) is swallowed; any other
    failure sets a flag the download button area checks and surfaces as
    the standard retry warning instead of silently succeeding."""
    if not _project_is_free_tier():
        return
    from sqlalchemy.exc import IntegrityError
    with db.get_session() as s:
        s.add(db.ArtifactEvent(
            user_id=current_user.id, project_key=_TRIAL_ARTIFACT_EVENT_SENTINEL_KEY,
            artifact_type=artifact_type,
        ))
        try:
            s.commit()
        except IntegrityError:
            s.rollback()
        except Exception:
            s.rollback()
            st.session_state["_free_artifact_mark_error"] = artifact_type


def _free_artifact_download_blocked(artifact_type: str) -> bool:
    """The single check a download button should gate on: True means show
    the paywall message instead of the file. False for every paid/
    unlimited/non-SaaS project (unlimited re-downloads -- Part B2's "buying
    a bid unlocks everything" rule), and for a free-tier project's FIRST
    download of one of the three FREE_TIER_ARTIFACTS; True for a free-tier
    project's second+ download of one of those three, or for ANY download
    of an artifact not on the free list at all (Methodology/Program PPTX,
    filled schedules -- see FREE_TIER_ARTIFACTS)."""
    if not _project_is_free_tier():
        return False
    if artifact_type not in FREE_TIER_ARTIFACTS:
        return True
    return _free_artifact_already_downloaded(artifact_type)


def _project_passes_status() -> dict:
    """Wraps auth.project_passes_status() for the current project -- see
    that function's docstring for the returned shape. Returns the
    "has_passes: False" shape outright for UNLIMITED_ACCOUNTS, non-SaaS
    use, or a project with no identity yet, so callers never need to
    special-case those themselves before reading ['remaining']."""
    if not IS_SAAS_MODE or not current_user or _access.get("unlimited"):
        return {"has_passes": False, "purchased": 0, "used": 0, "remaining": 0}
    _key = _current_project_key()
    if not _key:
        return {"has_passes": False, "purchased": 0, "used": 0, "remaining": 0}
    return auth.project_passes_status(current_user, _key)


# ---------------------------------------------------------------------------
# Part C -- rename-confirm dialog for a paid project. See _current_project_key()
# above for why project_name/tender_name/client_name/brief-hash together ARE
# the billing identity: changing any of the three typed names (Project
# Setup fields commit live, with no separate save step -- see that tab's
# own comment) computes a DIFFERENT key, meaning the paid analysis stays
# attached to the OLD identity and the new one looks unpaid. This warns
# once per actual identity change on an already-paid project, rather than
# silently letting someone rename their way into what looks like a fresh,
# unpaid project.
# ---------------------------------------------------------------------------

def _maybe_snapshot_paid_identity() -> None:
    """Call on EVERY rerun of the Project Setup tab (not just right after a
    live analysis run) to keep track of the most recently seen PAID
    (subscription/credit-funded) project identity AND the three editable
    field values (project/tender/client name) it currently corresponds to.
    Deliberately does nothing for a trial-funded or not-yet-funded project
    -- renaming one of those has no paid analysis to strand, so there's
    nothing to warn about.

    Audit fix Part 3b: this used to be called ONLY from inside the Tender
    Analysis tab's click handler, right after a paid analysis run actually
    completed -- so the snapshot lived purely in st.session_state and never
    survived a page refresh, a fresh login, or even just closing and
    reopening the tab. A rename attempted in ANY of those situations found
    no snapshot to compare against and silently let the edit through with
    no warning, orphaning the paid analysis under the old identity. Calling
    this on every Project Setup rerun instead re-derives "is the CURRENTLY
    loaded project paid?" from the database every time, so it doesn't
    matter whether the paid identity was established this session or three
    weeks and two logins ago -- as long as the project's fields, as loaded,
    still match what's on record as paid, the snapshot is current. Because
    Streamlit's text_input widgets commit on every keystroke (see the
    Project Setup tab's own comment), by the time an EDIT'S rerun reaches
    this function the fields already reflect the NEW, not-yet-paid
    identity -- auth.project_funded_by() correctly finds nothing for it, so
    this function simply does NOT overwrite the snapshot that rerun, and
    the previous (still-accurate) snapshot survives for
    _pending_rename_confirmation() to compare against on that exact same
    rerun."""
    if not IS_SAAS_MODE or not current_user:
        return
    _key = _current_project_key()
    if not _key:
        return
    _funded = auth.project_funded_by(current_user, _key)
    if _funded in ("subscription", "credit"):
        st.session_state["_paid_identity_key"] = _key.lower()
        st.session_state["_paid_identity_fields"] = {
            "project_name": st.session_state.get("project_name", ""),
            "tender_name": st.session_state.get("tender_name", ""),
            "client_name": st.session_state.get("client_name", ""),
        }


def _pending_rename_confirmation() -> str | None:
    """Returns the NEW project key a rename dialog should warn about, or
    None if nothing needs confirming right now. True exactly when the
    account has a remembered PAID identity (see
    _maybe_snapshot_paid_identity above) that no longer matches the
    CURRENT identity, and this exact new identity hasn't already been
    acknowledged (see _acknowledge_rename below) -- so the same dialog
    doesn't re-appear on every rerun after the user's already clicked
    through it once."""
    if not IS_SAAS_MODE or not current_user:
        return None
    _prev_paid_key = st.session_state.get("_paid_identity_key")
    if not _prev_paid_key:
        return None
    _key = (_current_project_key() or "").lower()
    if not _key or _key == _prev_paid_key:
        return None
    if st.session_state.get("_rename_ack_key") == _key:
        return None
    return _key


def _acknowledge_rename(new_key: str) -> None:
    """Marks `new_key` as acknowledged so _pending_rename_confirmation()
    doesn't show the same dialog again for it -- a safety net for the rare
    case where _confirm_rename()'s migration below couldn't actually apply
    (new_key collides with a different, already-existing project's own
    billing history -- see auth.migrate_project_identity()'s defensive
    "never clobber someone else's row" check), where the identity mismatch
    would otherwise persist and the dialog would reappear every rerun."""
    st.session_state["_rename_ack_key"] = new_key


def _confirm_rename(new_key: str) -> None:
    """"Yes, rename" -- audit fix Part 3b: migrates the paid project's
    ProposalUsage/ProjectPasses/ArtifactEvent rows from the old identity to
    `new_key` (see auth.migrate_project_identity()) so renaming a paid
    project never actually costs anyone their payment; the dialog's copy
    reflects this ("this project's payment will follow the new name").
    Re-snapshots the paid identity afterward (under whatever project_key is
    now current) so a FURTHER rename from here is tracked correctly too."""
    _old_key = st.session_state.get("_paid_identity_key")
    if _old_key:
        auth.migrate_project_identity(current_user, _old_key, new_key)
    _acknowledge_rename(new_key)
    _maybe_snapshot_paid_identity()


def _cancel_rename(new_key: str) -> None:
    """"Cancel" -- audit fix Part 3b: actually cancels, by restoring the
    three editable identity fields to what they were the last time this
    project's identity was confirmed paid (see
    _maybe_snapshot_paid_identity's _paid_identity_fields snapshot).
    Previously both dialog buttons called the same _acknowledge_rename(),
    which only dismissed the dialog -- the typed field kept showing
    whatever was just typed either way, so "Cancel" didn't actually cancel
    the edit. Reassigning a widget-bound session_state key here (a
    callback, called before the next rerun re-instantiates the text_input)
    is exactly the supported way to programmatically revert one; doing the
    same assignment inline, after the widget has already rendered in the
    SAME script run, is what Streamlit disallows -- this runs from the
    dialog button's own click handling, one rerun later, which is fine."""
    _fields = st.session_state.get("_paid_identity_fields") or {}
    for _field, _value in _fields.items():
        st.session_state[_field] = _value
    _acknowledge_rename(new_key)


def _record_ai_click() -> None:
    """Call once, right when an AI-feature button's own click-handler
    begins real work -- the first line inside its `try:`/spinner block, not
    inside the `ready`/disabled computation (which reruns on every
    interaction across the whole app, not just AI-feature clicks). Advances
    the fair-use rate limit's rolling-window counter (modules/limits.py) for
    this account; the `ready` boolean already refused the click if the
    account was over its cap (see _ai_gate_msg, a read-only peek computed
    once per script run in 00_init.py), so this never needs to abort
    anything itself -- it only affects whether the NEXT click is allowed.
    Fire-and-forget; never raises; a no-op outside SaaS mode."""
    if IS_SAAS_MODE and current_user:
        try:
            limits.record_ai_call(current_user.id, not limits.is_paid_tier(_access))
        except Exception:
            pass


def _company_materials_flags() -> dict:
    flags = {f"has_{k}": bool(st.session_state.company_material_text.get(k)) for k in COMPANY_MATERIAL_CATEGORIES}
    flags["has_project_photos"] = bool(st.session_state.project_photo_bytes)
    flags["has_company_image_library"] = bool(st.session_state.branding_bytes)
    # Firm-profile facts count as materials too -- an insurance requirement
    # is covered when the firm profile actually holds the insurance, not only
    # when a company-profile document happens to have been uploaded.
    flags.update(_firm_materials_flags())
    return flags


def _rebuild_structure():
    """Re-derive weighting -> pages -> sections -> guidance notes as one consistent chain.
    proposal_format only changes which branch build_proposal_structure() takes at the end --
    weighting and page allocation are always derived the same way from the same brief."""
    analysis = st.session_state.analysis
    weighted = weighting_engine.apply_weighting(analysis)
    allocations = page_allocation.allocate_pages(weighted, analysis)
    sections = proposal_structure.build_proposal_structure(
        analysis, weighted, allocations, proposal_format=st.session_state.proposal_format,
    )
    st.session_state.weighted_criteria = weighted
    st.session_state.allocations = allocations
    st.session_state.sections = sections
    st.session_state.guidance_notes = guidance_generator.generate_all_guidance_notes(sections)
    # Remember which format these sections were built for -- see
    # _structure_format_stale() below. Without this, switching the Proposal
    # format selector (tab 1) after already generating structure leaves
    # st.session_state.sections holding the OLD format's section titles
    # (e.g. "Executive Summary"/"Relevant Experience" instead of "Project
    # Understanding"/"Methodology and Deliverables"), so _draftable_sections()
    # silently matches nothing -- drafting "succeeds" against an empty list
    # and the user sees "Draft generation complete" with nothing to show for
    # it, no error anywhere.
    st.session_state["_sections_built_format"] = st.session_state.proposal_format


def _reset_downstream_from_brief() -> None:
    """Resets every piece of state derived from the tender brief -- call this
    whenever the brief itself is replaced or cleared (see the Upload Docs
    tab: a fresh upload that changes the file signature, or the "Clear all"
    button). Without this, replacing the brief only reset tender_extracted
    itself, leaving Tender Analysis, Structure, Page Allocation, Draft
    Responses, Graphics & Design, Team & Resourcing, and Fee Estimate all
    still holding the PREVIOUS brief's results -- so the sidebar stepper
    kept showing those steps as done (green) for a brand new project that
    hadn't actually gone through them yet. Deliberately leaves Project
    Setup fields (project/client/tender name, etc.) and firm-level company
    materials (CV library, past proposals, branding) untouched -- neither
    is derived from the brief itself, so there's no reason to clear them
    when the brief changes."""
    downstream_defaults = {
        "analysis": None,
        "weighted_criteria": None,
        "allocations": None,
        "sections": None,
        "guidance_notes": None,
        "compliance_items": None,
        "gap_items": None,
        "drafts": None,
        "executive_summary": None,
        "team_intro": None,
        "experience_intro": None,
        "pitch_review": None,
        "pitch_questions": None,
        "tender_summary_buffer": None,
        "graphics": None,
        "weighting_chart_png": None,
        "fee_estimates": None,
        "docx_buffer": None,
        "divider_images": {},
        "cover_hero_png": None,
        "_sections_built_format": None,
        "resource_plan": [],
        "discipline_fee_lines": [],
        "scope_item_fees": [],
        "org_chart_png": None,
        "reference_projects": None,
        "reference_project_photos": {},
        "reference_project_warnings": [],
        "program_schedule": {},
        "program_week_labels": [],
        "personnel_photos": {},
        "personnel_inclusion_suggestions": {},
        "dismissed_disciplines": [],
        "dismissed_fee_disciplines": [],
        "fee_seed_total": 0.0,
        "fee_estimate_manual_total": 0.0,
        "letter_fee_total_override": 0.0,
    }
    for key, value in downstream_defaults.items():
        st.session_state[key] = value


def _structure_format_stale() -> bool:
    """True when sections exist but were generated under a different Proposal
    format than the one currently selected -- see the comment in
    _rebuild_structure() above for why this matters."""
    return (
        st.session_state.sections is not None
        and st.session_state.get("_sections_built_format") != st.session_state.proposal_format
    )


def _is_letter() -> bool:
    return st.session_state.proposal_format == "letter"


def _firm_profile() -> "object | None":
    """This account's firm profile, or None.

    Cached for the run: it is read by the sidebar, the project seeder and
    several exporters, and re-querying per rerun for a row that changes
    perhaps once a year is waste. Never raises -- a database hiccup must
    degrade to "no profile" (i.e. today's placeholders), not to an error
    page in the middle of a bid."""
    if "_firm_profile_cache" not in st.session_state:
        try:
            st.session_state._firm_profile_cache = firm_profile.get_profile(
                current_user.id if (IS_SAAS_MODE and current_user) else None  # noqa: F821
            )
        except Exception:
            st.session_state._firm_profile_cache = None
    return st.session_state._firm_profile_cache


def _firm_profile_is_empty() -> bool:
    return firm_profile.is_empty(_firm_profile())


def _structured_material_by_section(sections: list) -> dict:
    """{section_title: material} for sections whose content the user has
    already reviewed in structured form.

    Relevant Experience is drafted from the reference-project cards the user
    edited, and Key Personnel from the resourcing plan's own profiles,
    instead of from the raw truncated text of whatever was uploaded. Those
    two disagreed inside the same exported document: the cards showed the
    user's corrections and the drafted prose argued from the original."""
    out = {}
    for section in sections or []:
        title = getattr(section, "title", "")
        lowered = title.lower()
        if "experience" in lowered and st.session_state.reference_projects:
            lines = []
            for ref in st.session_state.reference_projects:
                lines.append(f"--- Reference project: {ref.title} ---")
                if ref.client:
                    lines.append(f"Client: {ref.client}")
                if ref.description:
                    lines.append(ref.description)
                if ref.relevance_text:
                    lines.append(f"Why it is relevant: {ref.relevance_text}")
                if ref.personnel_involved:
                    lines.append("Personnel involved: " + ", ".join(ref.personnel_involved))
            out[title] = (
                "--- USER-REVIEWED REFERENCE PROJECTS (these are the corrected, "
                "structured entries -- use ONLY these, not any raw uploaded text) ---\n"
                + "\n".join(lines)
            )
        elif any(word in lowered for word in ("risk", "safety", "quality")) and st.session_state.risk_register:
            block = risk_register.format_for_prompt(st.session_state.risk_register)
            if block:
                out[title] = block
        elif "personnel" in lowered and st.session_state.resource_plan:
            from modules.resourcing import personnel_profiles_deduped
            lines = []
            for entry in personnel_profiles_deduped(st.session_state.resource_plan):
                # The include/exclude tick lives on the assignment itself, not
                # on the flattened profile dict -- someone unticked must not
                # be named in the drafted prose either.
                assignment = entry.get("assignment")
                if not getattr(assignment, "include_in_proposal", True):
                    continue
                if not (entry.get("name") or "").strip():
                    continue
                lines.append(f"--- {entry['name']} ({', '.join(entry.get('roles') or [])}) ---")
                for label, key in (("Qualification", "qualification"), ("RPEQ", "rpeq_status"),
                                   ("Years of experience", "years_experience"),
                                   ("On this project they will", "value_to_project")):
                    value = (entry.get(key) or "").strip()
                    if value:
                        lines.append(f"{label}: {value}")
                for project in entry.get("relevant_projects") or []:
                    lines.append(f"Relevant project: {project}")
            if lines:
                out[title] = (
                    "--- NOMINATED PERSONNEL, AS ENTERED BY THE BID TEAM (use ONLY these "
                    "people and only these facts about them) ---\n" + "\n".join(lines)
                )
    return out


def _firm_materials_flags() -> dict:
    """Firm-profile facts the compliance matrix should treat as covered."""
    profile = _firm_profile()
    return {
        "firm_profile_has_insurances": bool(firm_profile.insurances(profile)),
        "firm_profile_has_certifications": bool(firm_profile.certifications(profile)),
    }


def _firm_export_context() -> dict:
    """The firm profile as the exporters want it (see
    firm_profile.export_context) -- one bundle rather than six new keyword
    arguments on two already-long builder signatures."""
    return firm_profile.export_context(_firm_profile(), st.session_state.get("bidder_name", ""))


def _seed_project_from_firm_profile() -> None:
    """Fill this project's blank fields from the firm profile.

    Only ever writes where the project's own value is still empty. A seed
    that overwrote something typed for THIS bid would be the worst kind of
    bug in a proposal tool -- silent, and wrong in a document."""
    profile = _firm_profile()
    if profile is None:
        return
    for key, value in firm_profile.project_seed(profile).items():
        if not (st.session_state.get(key) or "").strip():
            st.session_state[key] = value
    # A firm that has recorded its QA/Work Verification commitment has
    # answered the methodology table's WVR question -- see Batch 2's honesty
    # fix. Still only a seed: unticking it on a bid sticks for that bid.
    if (getattr(profile, "qa_statement", "") or "").strip():
        st.session_state.methodology_wvr_confirmed = True


FEE_PRESENTATION_LABELS = {
    "pct_split": "% split by discipline",
    "discipline_buildup": "discipline build-up",
    "scope_buildup": "scope-item build-up",
}
# Audit fix Part 6: i18n.t() keys for the same three labels, used wherever
# they're shown to the user (_fee_inclusion_summary) -- FEE_PRESENTATION_LABELS
# itself stays English-only since it's also used as plain internal dict keys
# in a couple of places that don't want translated text.
FEE_PRESENTATION_LABEL_KEYS = {
    "pct_split": "fee_label_pct_split",
    "discipline_buildup": "fee_label_discipline_buildup",
    "scope_buildup": "fee_label_scope_buildup",
}


def _fee_presentation_has_data(key: str) -> bool:
    if key == "pct_split":
        return bool(st.session_state.get("fee_estimates"))
    if key == "discipline_buildup":
        return any(
            (l.total_hours or l.rate_per_hour)
            for l in (st.session_state.get("discipline_fee_lines") or [])
        )
    return any(
        (getattr(f, "fee_amount", 0) or 0)
        for f in (st.session_state.get("scope_item_fees") or [])
    )


def _fee_include_checkbox(key: str, widget_key: str) -> None:
    """The "include this fee presentation in the proposal" tick.

    The app builds up to three parallel fee presentations and which of them
    reached the proposal used to be hardcoded per pack format -- so a user
    who priced a discipline build-up for an audience that wanted a
    percentage split had no way to say so."""
    included = st.session_state.fee_sections_included
    has_data = _fee_presentation_has_data(key)
    ticked = st.checkbox(
        i18n.t("fee_include_presentation_checkbox"),
        value=bool(included.get(key)), key=widget_key, disabled=not has_data,
    )
    if has_data:
        included[key] = ticked
    else:
        st.caption(i18n.t("fee_enter_figures_first_caption"))


def _fee_inclusion_summary() -> None:
    """Standing line under the Fee tab header naming what will be exported."""
    included = st.session_state.fee_sections_included
    chosen = [i18n.t(FEE_PRESENTATION_LABEL_KEYS[k]) for k in ("pct_split", "discipline_buildup", "scope_buildup")
              if included.get(k)]
    if chosen:
        st.caption(i18n.t("fee_included_in_proposal_caption") + " " + " + ".join(chosen))
    else:
        st.warning(i18n.t("fee_none_ticked_warning"))


def _program_model_from_state():
    """The one program model every output renders from -- preview, letter
    pack, and companion deck. Built here rather than three times over so a
    preview and an export cannot disagree."""
    return program_render.build_model(
        st.session_state.get("program_schedule") or {},
        st.session_state.get("program_week_labels") or [],
        st.session_state.get("methodology_stages") or [],
        st.session_state.get("program_start_date"),
        st.session_state.get("analysis"),
        st.session_state.get("project_name") or "",
        st.session_state.get("client_name") or "",
    )


def _program_signature(style: str) -> tuple:
    """Everything the drawn program depends on. Anything not in here is
    something a change to could leave a stale picture on screen.

    Also reused (see 80_export.py) as the cache signature for the exported
    delivery-program PPTX, which DOES vary with output_language now (Audit
    Round 2, Part 5) -- included here so a language switch busts that cache
    instead of serving a stale-language deck."""
    stages = st.session_state.get("methodology_stages") or []
    start = st.session_state.get("program_start_date")
    analysis = st.session_state.get("analysis")
    return (
        style,
        tuple((title, tuple(bool(w) for w in (weeks or [])))
              for title, weeks in (st.session_state.get("program_schedule") or {}).items()),
        tuple(st.session_state.get("program_week_labels") or []),
        tuple((getattr(s, "name", ""), getattr(s, "week_start", None), getattr(s, "week_end", None),
               tuple(getattr(s, "engagement_activities", None) or [])) for s in stages),
        start.isoformat() if hasattr(start, "isoformat") else None,
        (getattr(analysis, "submission_date", "") or "") if analysis is not None else "",
        st.session_state.get("project_name") or "",
        st.session_state.get("client_name") or "",
        st.session_state.get("proposal_theme") or "",
        st.session_state.get("output_language") or "en",
    )


# Rendering a program PNG is ~0.4s of matplotlib. The preview redraws on
# every Streamlit rerun -- which is every keystroke in the fee tables on the
# same tab -- so it is cached on the content it was drawn from, not on time.
_PROGRAM_PNG_CACHE_SIZE = 8


def _program_png(style: str | None = None) -> bytes | None:
    """The current program drawn in `style`, cached on its content signature.
    None when it could not be drawn; every caller has a fallback."""
    style = style or st.session_state.get("program_style") or program_render.DEFAULT_STYLE
    signature = _program_signature(style)
    cache = st.session_state.setdefault("_program_png_cache", {})
    if signature in cache:
        return cache[signature]
    try:
        accent = f"#{export_docx._theme_colours(st.session_state.get('proposal_theme'))['accent']}"
        png = program_render.render_png(_program_model_from_state(), style, accent)
    except Exception as exc:  # noqa: BLE001 -- a preview is never worth a traceback
        print(f"[program preview] {exc}", file=sys.stderr)
        png = None
    if len(cache) >= _PROGRAM_PNG_CACHE_SIZE:
        cache.pop(next(iter(cache)))
    cache[signature] = png
    return png


def _program_style_control(key_suffix: str) -> None:
    """The four-way program presentation style picker, plus a live preview
    drawn from this project's own data.

    A preview rather than four stock thumbnails: the styles differ most in
    how they cope with THIS program's activity names and week count, which a
    generic sample cannot show."""
    st.markdown("**Program presentation style**")
    st.caption(
        "How the program is drawn in the proposal and the companion PowerPoint. "
        "Purely presentational -- it never changes a single week you ticked."
    )
    st.radio(
        "Program presentation style", program_render.STYLES,
        key="program_style", horizontal=True, label_visibility="collapsed",
        format_func=lambda s: program_render.STYLE_LABELS.get(s, s),
    )
    chosen = st.session_state.program_style
    st.caption(program_render.STYLE_DESCRIPTIONS.get(chosen, ""))

    model = _program_model_from_state()
    if model.is_empty:
        st.info("Build the delivery program above and the preview appears here.")
        return
    effective = program_render.effective_style(model, chosen)
    if effective != chosen:
        st.info(
            "**Stage grouping needs methodology stages** -- generate them in the Drafting "
            "step and this becomes swimlanes. Showing the Gantt style for now."
        )
    png = _program_png(chosen)
    if png:
        st.image(png, use_container_width=True)
        st.caption("This is exactly what the exported pack and the PowerPoint will show.")
    else:
        st.caption(
            "Couldn't draw the preview just now -- the program itself is unaffected, and the "
            "export falls back to a plain week-by-week table."
        )


def _fee_apply_control(state_prefix: str, pending: bool, indent_note: str = "totals") -> bool:
    """The explicit "Apply changes" control under a deferred-apply fee table.

    Replaces a checkbox labelled "Done entering data -- refresh totals". A
    checkbox that acts as a button reads, to anyone who hasn't been told
    otherwise, as a statement about the data -- so an unticked box next to a
    total looked like the app saying the total was wrong, and users ticked
    and unticked it trying to work out what it meant. A primary button says
    what it does.

    The deferred-apply behaviour itself is unchanged and deliberate: these
    tables live in fragments so that typing in them doesn't rerun the whole
    script, and applying on every keystroke would fight the editor's own
    commit timing (see the race-protection notes at each call site). What
    changes is only how the user is asked to apply, plus an explicit warning
    while an edit is outstanding -- previously the only signal was a total
    that quietly disagreed with the table above it.
    """
    if pending:
        st.warning(i18n.t("fee_table_edited_warning", indent_note=indent_note))
    return st.button(
        i18n.t("fee_apply_changes_button"), type="primary", key=f"{state_prefix}apply_btn",
        disabled=not pending,
        help=None if pending else i18n.t("fee_nothing_to_apply_help"),
    )


# Below this a "draft" is not a draft. An AI call that returns a sentence
# and a heading has failed at the job, but it looks like success from the
# outside -- the expander opens, there is text in it.
MIN_USEFUL_DRAFT_WORDS = 40


def _thin_drafts(drafts: dict) -> list[str]:
    """Section titles whose draft came back empty or barely there.

    A blank draft rendered as a blank expander under a green "Draft
    generation complete", which is the app asserting success about a thing
    that did not happen. Nobody scrolls twelve expanders checking."""
    thin = []
    for title, draft in (drafts or {}).items():
        text = (getattr(draft, "draft_text", "") or "").strip()
        if not text:
            thin.append(f"{title} (empty)")
        elif len(text.split()) < MIN_USEFUL_DRAFT_WORDS:
            thin.append(f"{title} ({len(text.split())} words)")
    return thin


def _placeholders_in_generated_pack(buffer) -> list:
    """Reads a just-generated proposal DOCX back and lists the placeholders
    really in it, for the Tender Summary's User Input Required list. Never
    raises -- a failed sweep must not block the export it is describing."""
    try:
        import io as _io

        from docx import Document as _Document

        from modules.export_docx import collect_placeholders
        return collect_placeholders(_Document(_io.BytesIO(buffer.getvalue())))
    except Exception:
        return []


def _export_input_signature() -> str:
    """A hash of everything that ends up in the exported pack.

    Stamped when a pack is generated and compared on every render, so a
    download button can say the file predates the user's latest edits.
    Before this, editing a fee after generating the DOCX silently served the
    OLD pack -- the worst possible failure mode here, because the user gets
    a plausible document that is quietly wrong, and nothing on screen
    disagrees with them."""
    import hashlib
    import json as _json

    def _dump(value):
        try:
            if hasattr(value, "model_dump"):
                return value.model_dump()
            if isinstance(value, dict):
                return {str(k): _dump(v) for k, v in value.items()}
            if isinstance(value, (list, tuple)):
                return [_dump(v) for v in value]
            if isinstance(value, (str, int, float, bool)) or value is None:
                return value
            if isinstance(value, (bytes, bytearray)):
                # Bytes are photos/charts -- length is enough to notice a swap
                # without hashing megabytes on every rerun.
                return f"bytes:{len(value)}"
            return str(value)
        except Exception:
            return "?"

    keys = [
        "project_name", "client_name", "tender_name", "bidder_name", "submission_date_input",
        "proposal_theme", "project_type", "proposal_format", "body_font",
        "analysis", "sections", "drafts", "guidance_notes", "compliance_items", "gap_items",
        "graphics", "fee_estimates", "discipline_fee_lines", "scope_item_fees",
        "resource_plan", "reference_projects", "executive_summary", "team_intro",
        "experience_intro", "methodology_stages", "methodology_wvr_confirmed", "risk_register",
        "program_schedule", "program_week_labels", "terms_of_engagement_text",
        "project_differentiator", "project_sales_pitch", "fee_estimate_manual_total",
        "cover_photo_index", "fee_sections_included", "program_style", "program_start_date",
        "org_chart_style",
        "letter_sender_name", "letter_sender_title",
        "letter_sender_phone", "letter_sender_email", "letter_sender_address",
    ]
    payload = {key: _dump(st.session_state.get(key)) for key in keys}
    payload["_photos"] = [len(b or b"") for b in (st.session_state.get("project_photo_bytes") or [])]
    payload["_org_chart"] = len(st.session_state.get("org_chart_png") or b"")
    try:
        blob = _json.dumps(payload, sort_keys=True, default=str)
    except Exception:
        blob = str(payload)
    return hashlib.sha256(blob.encode("utf-8", "replace")).hexdigest()


# ---------------------------------------------------------------------------
# Audit fix Part 3a -- pass metering for a full "Generate/Regenerate All
# Drafts" run on a PAID project. Part B2 defines a pass as one full
# generation cycle: the initial Tender Analysis run, OR a later
# regeneration once a tracked input has actually changed. Tender Analysis's
# own repeat-run metering lives in the Tender Analysis tab (see
# auth.consume_project_pass()'s call site there); this is the matching
# metering for the OTHER full-cycle action in the app, "Generate/Regenerate
# All Drafts" (Draft Responses tab) -- previously gated only on
# _current_project_already_paid(), with no pass cost at all, so a paid
# project could regenerate every section's drafts an unlimited number of
# times regardless of its 5-pass allowance.
# ---------------------------------------------------------------------------

def _draft_generation_input_signature() -> str:
    """Same hashing technique as _export_input_signature() above, but
    scoped to what actually FEEDS a full draft-generation run, not the
    whole export bundle -- deliberately excludes "drafts" itself (and
    other pure OUTPUTS of this same button, like executive_summary/
    team_intro/pitch_review) because including them would make the
    signature circular: every full regenerate overwrites drafts, so
    comparing a "before" and "after" signature that both include drafts
    would always disagree, defeating the entire point of "unchanged
    inputs, no pass spent." What IS compared is everything that actually
    changes what the drafter would produce: the brief's own analysis, the
    built section list, compliance items, the resourcing/team plan,
    reference projects, the differentiator/sales-pitch text, the raw
    company material, and the chosen output language."""
    import hashlib
    import json as _json

    def _dump(value):
        try:
            if hasattr(value, "model_dump"):
                return value.model_dump()
            if isinstance(value, dict):
                return {str(k): _dump(v) for k, v in value.items()}
            if isinstance(value, (list, tuple)):
                return [_dump(v) for v in value]
            if isinstance(value, (str, int, float, bool)) or value is None:
                return value
            if isinstance(value, (bytes, bytearray)):
                return f"bytes:{len(value)}"
            return str(value)
        except Exception:
            return "?"

    keys = [
        "analysis", "sections", "compliance_items", "resource_plan",
        "reference_projects", "project_differentiator", "project_sales_pitch",
        "company_material_text", "output_language", "proposal_format",
    ]
    payload = {key: _dump(st.session_state.get(key)) for key in keys}
    try:
        blob = _json.dumps(payload, sort_keys=True, default=str)
    except Exception:
        blob = str(payload)
    return hashlib.sha256(blob.encode("utf-8", "replace")).hexdigest()


def _draft_project_is_metered() -> tuple[bool, str]:
    """(is_metered, project_key) -- True only for a PAID (subscription/
    credit) project: never for a trial-funded project (which gets exactly
    ONE generation pass total, already enforced at the Tender Analysis
    level -- Part B's original one-pass-free-tier promise includes a full
    drafted pack from that one pass, so drafting itself stays unmetered
    for a trial project), never for UNLIMITED_ACCOUNTS, and never outside
    SaaS mode."""
    if not IS_SAAS_MODE or not current_user or _access.get("unlimited"):
        return False, ""
    _key = _current_project_key()
    if not _key or auth.is_trial_funded_project(current_user, _key):
        return False, ""
    return True, _key


def _draft_would_consume_pass() -> bool:
    """True if clicking 'Generate/Regenerate All Drafts' right now would
    spend one of a paid project's passes -- i.e. this project is paid (see
    _draft_project_is_metered()) AND its drafting inputs have changed
    since the last metered run. False for the very first draft generation
    on a freshly-analysed paid project (no baseline signature recorded
    yet -- that generation is covered by the pass Tender Analysis itself
    already spent, same as re-downloading unchanged documents is free).
    Read-only and side-effect-free; used for both the pre-click caption
    and the click handler itself, so they can never disagree."""
    _metered, _key = _draft_project_is_metered()
    if not _metered:
        return False
    _last_sig = st.session_state.get("_last_draft_metered_signature")
    if _last_sig is None:
        return False
    return _draft_generation_input_signature() != _last_sig


def _generated_language_stale() -> bool:
    """Audit fix Part 5: True when drafts exist but were generated in a
    DIFFERENT language than the project's current output_language -- e.g.
    the user drafted in English, then switched the project to Spanish
    without regenerating. Purely informational (drives a non-blocking
    notice on the drafting/export tabs); never triggers a regeneration
    itself. False before any draft has ever been generated (generated_
    language is None) -- nothing to warn about yet."""
    _generated = st.session_state.get("generated_language")
    if not _generated:
        return False
    return _generated != st.session_state.get("output_language", "en")


def _mark_export_generated() -> None:
    st.session_state._export_signature = _export_input_signature()
    # Every generated pack leaves its priced split behind, so the next bid of
    # this type can start from what this firm actually charged rather than
    # from a published average. Best-effort by design: a bookkeeping failure
    # must never cost someone the pack they just waited for.
    _record_fee_snapshot()


def _export_is_stale() -> bool:
    stamped = st.session_state.get("_export_signature")
    if not stamped:
        return False
    return stamped != _export_input_signature()


def _export_readiness() -> list[dict]:
    """What's still outstanding before this pack is worth sending.

    Most of the red placeholder walls in an exported pack are not missing
    information -- they are a step the user hasn't run yet. Listing them
    here, each with where to go, turns a silently red document into a short
    list of actions."""
    items: list[dict] = []

    def add(label: str, where: str, detail: str = "") -> None:
        items.append({"label": label, "where": where, "detail": detail})

    sections = st.session_state.sections or []
    drafts = st.session_state.drafts or {}
    undrafted = [
        s.title for s in _draftable_sections(sections)
        if not (drafts.get(s.title) and (drafts[s.title].draft_text or "").strip())
    ]
    if undrafted:
        add(f"{len(undrafted)} section(s) have no draft", "Draft Responses",
            ", ".join(undrafted[:4]) + (", ..." if len(undrafted) > 4 else ""))

    # Drafts that came back far off their page budget.
    from modules.draft_generator import length_verdict
    off_budget = [
        f"{s.title} ({length_verdict(s, drafts[s.title])})"
        for s in sections
        if s.title in drafts and length_verdict(s, drafts[s.title])
    ]
    if off_budget:
        add(f"{len(off_budget)} draft(s) are well off their page budget", "Draft Responses",
            ", ".join(off_budget[:4]))

    plan = st.session_state.resource_plan or []
    unassigned = [a.slot for a in plan if not (a.person_name or "").strip()]
    if unassigned:
        add(f"{len(unassigned)} role(s) have nobody assigned", "Team & Resourcing",
            ", ".join(unassigned[:4]) + (", ..." if len(unassigned) > 4 else ""))
    if not plan:
        add("No team assigned", "Team & Resourcing", "Key Personnel will export empty.")

    fee_lines = st.session_state.discipline_fee_lines or []
    unpriced = [l.discipline for l in fee_lines if not (l.total_hours or l.rate_per_hour)]
    if unpriced:
        add(f"{len(unpriced)} discipline(s) are unpriced", "Fees & Program",
            ", ".join(unpriced[:4]) + (", ..." if len(unpriced) > 4 else ""))
    if not fee_lines:
        add("No fee build-up entered", "Fees & Program", "The fee table exports as a placeholder.")

    included = st.session_state.get("fee_sections_included") or {}
    if not any(included.values()):
        add("No fee presentation ticked", "Fees & Program",
            "The proposal's fee section will export as a red placeholder.")
    else:
        zero_rows = []
        if included.get("discipline_buildup"):
            zero_rows += [l.discipline for l in (st.session_state.discipline_fee_lines or [])
                          if not (l.total_hours or l.rate_per_hour)]
        if included.get("scope_buildup"):
            zero_rows += [f.item_title for f in (st.session_state.scope_item_fees or [])
                          if not (getattr(f, "fee_amount", 0) or 0)]
        if zero_rows:
            add(f"{len(zero_rows)} zero-value row(s) in a fee table you're exporting",
                "Fees & Program",
                ", ".join(zero_rows[:4]) + (", ..." if len(zero_rows) > 4 else ""))

    if not st.session_state.program_schedule:
        add("No delivery program", "Fees & Program",
            "Also blocks the derived cash-flow profile.")
    elif not st.session_state.get("program_start_date"):
        add("No anticipated start date", "Fees & Program",
            "Optional -- set it and the program shows real dates instead of week numbers.")

    if st.session_state.methodology_stages is None and not _is_letter():
        add("Design stages not drafted", "Draft Responses",
            "The methodology table exports with placeholder columns.")
    if st.session_state.risk_register is None and (getattr(st.session_state.analysis, "risks", None)):
        add("Risk register not drafted", "Draft Responses",
            "The brief's risks export as raw bullets.")

    profile = _firm_profile()
    if firm_profile.is_empty(profile):
        add("Firm profile is empty", "Sidebar -> Firm profile",
            "About ten red placeholders in the pack come from here.")
    else:
        missing = []
        if not (getattr(profile, "logo_bytes", None)):
            missing.append("logo")
        if not (getattr(profile, "abn", "") or "").strip():
            missing.append("ABN")
        if not (getattr(profile, "registered_address", "") or "").strip():
            missing.append("registered address")
        if not firm_profile.insurances(profile):
            missing.append("insurances")
        if missing:
            add("Firm profile is incomplete", "Sidebar -> Firm profile", ", ".join(missing))

    if not st.session_state.executive_summary:
        add("No executive summary", "Draft Responses", "The pack's first page exports empty.")

    return items


def photo_key_for(obj, fallback: str) -> str:
    """The key an object's photo is stored under: its stable id, or its
    name/title for anything minted before ids existed."""
    return (getattr(obj, "photo_id", "") or "").strip() or (fallback or "").strip()


def _ensure_photo_ids() -> None:
    """Give every person and reference project a stable photo id, migrating
    any photo currently filed under their name.

    Both photo dicts were keyed by a name the user can edit, so fixing a
    typo in "Mat Willliams" orphaned his headshot: it stayed in the dict
    under the old spelling and just stopped appearing, with nothing to
    explain why. Ids are minted here on first sight -- including for
    projects loaded from a file saved before this existed -- and the photo
    moves across at the same moment, so the migration happens once and
    invisibly."""
    import uuid

    for assignment in (st.session_state.get("resource_plan") or []):
        if getattr(assignment, "photo_id", ""):
            continue
        assignment.photo_id = uuid.uuid4().hex[:12]
        name = (getattr(assignment, "person_name", "") or "").strip()
        photos = st.session_state.get("personnel_photos") or {}
        if name and name in photos:
            photos[assignment.photo_id] = photos[name]

    for project in (st.session_state.get("reference_projects") or []):
        if getattr(project, "photo_id", ""):
            continue
        project.photo_id = uuid.uuid4().hex[:12]
        title = (getattr(project, "title", "") or "").strip()
        photos = st.session_state.get("reference_project_photos") or {}
        if title and title in photos:
            photos[project.photo_id] = photos[title]


def _cover_photo_bytes():
    """The photo chosen for the cover, or the first uploaded one.

    Was always photos[0] with no way to change it, so the single most
    visible image in the pack was decided by upload order."""
    photos = st.session_state.get("project_photo_bytes") or []
    if not photos:
        return None
    index = st.session_state.get("cover_photo_index") or 0
    return photos[index] if 0 <= index < len(photos) else photos[0]


def _dates_look_equivalent(a: str, b: str) -> bool:
    """Loose comparison of two human-written dates.

    Deliberately forgiving: "14 July 2026", "14/07/2026" and "14 Jul 2026"
    are the same date written three ways, and warning about those would
    train the user to ignore the warning. Only a genuine difference in the
    day, month or year should fire it."""
    import re as _re

    _MONTHS = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }

    def _parts(text: str):
        text = (text or "").lower()
        year = None
        year_match = _re.search(r"\b(20\d{2})\b", text)
        if year_match:
            year = int(year_match.group(1))
        month = None
        for name, number in _MONTHS.items():
            if name in text:
                month = number
                break
        numbers = [int(n) for n in _re.findall(r"\b(\d{1,2})\b", text)]
        day = numbers[0] if numbers else None
        if month is None and len(numbers) >= 2:
            # A numeric date: assume day/month, the Australian convention
            # this app is written for.
            day, month = numbers[0], numbers[1]
        return day, month, year

    da, ma, ya = _parts(a)
    db, mb, yb = _parts(b)
    if not any((da, ma, ya)) or not any((db, mb, yb)):
        # One of them isn't recognisably a date at all -- say nothing rather
        # than warn about free text.
        return True
    for left, right in ((da, db), (ma, mb), (ya, yb)):
        if left is not None and right is not None and left != right:
            return False
    return True


def _program_week_count() -> int:
    """How many week columns the delivery program actually has.

    Read from the schedule itself rather than from program_num_weeks: the two
    drift apart the moment someone changes the week count without pressing
    "Generate default program" again, and week labels that outnumber the
    grid's columns (or fall short of them) silently mislabel the whole
    program."""
    rows = (st.session_state.get("program_schedule") or {}).values()
    widths = [len(row) for row in rows]
    return max(widths) if widths else int(st.session_state.get("program_num_weeks") or 0)


# Small Scope pack sections that are actually free text to be AI-drafted-then-edited, same as
# every section in the Large Scope pack. The other sections (Scope of Work, Project Team, Fees,
# Program) are built from structured data the user supplies directly (scope_items already
# extracted from the brief, CV-drafted team bios, the fee table, the program grid) -- running
# them through the generic AI drafter would just waste calls on content nothing downstream uses.
LETTER_DRAFTABLE_TITLES = {"Project Understanding", "Methodology and Deliverables"}


def _draftable_sections(sections: list) -> list:
    if st.session_state.proposal_format == "letter":
        return [s for s in sections if s.title in LETTER_DRAFTABLE_TITLES]
    return sections


def _project_identifier() -> str:
    """What a local auto-save file / "Recent projects" entry gets named after.
    See local_project_store.project_identifier() -- prefers the descriptive
    project name over the tender/EOI name, which is often left generic."""
    return local_project_store.project_identifier(st.session_state.project_name, st.session_state.tender_name)


def _files_signature(files) -> tuple:
    """A cheap identity for a set of uploaded files (name + size), so we can tell
    when the uploads have actually changed and avoid re-extracting the same files
    on every Streamlit rerun (reruns fire on nearly every interaction, and
    re-parsing large PDFs each time is what makes the app feel frozen)."""
    return tuple((getattr(f, "name", ""), getattr(f, "size", None)) for f in (files or []))


# ---------------------------------------------------------------------------
# Fee prepopulation -- three labelled tiers, best first (Batch 10)
#
# The fee sheet used to open at zeros with one "benchmark" behind it: a
# bundled table of published industry averages. The ladder below adds the two
# tiers that are actually about THIS firm and THIS brief, and every figure
# any of them produces says which tier it came from, so a user can never mix
# up their own past pricing with a rule of thumb or with an AI's recollection.
# ---------------------------------------------------------------------------

def _fee_history_user_id() -> str | None:
    """The account fee snapshots belong to. None in local mode, where there
    is no user to scope a history to -- history is per firm, and a shared
    local install has no way to tell whose pricing is whose."""
    try:
        return current_user.id if (IS_SAAS_MODE and current_user) else None  # noqa: F821
    except Exception:  # noqa: BLE001
        return None


def _record_fee_snapshot() -> None:
    """Leave this project's priced split behind for the next bid of the same
    type. Called when a pack is generated and when one is archived; the
    upsert in fee_history keeps that from counting as two bids."""
    fee_history.record_snapshot(
        _fee_history_user_id(),
        _current_project_key(),
        st.session_state.get("project_type") or "",
        st.session_state.get("discipline_fee_lines") or [],
        st.session_state.get("project_name") or "",
    )


def _firm_rate_card() -> dict:
    """The firm's standard charge-out rates, or {} -- including when the Firm
    Profile has never been filled in, which is the normal case for a new
    account and must not be an error.

    Cached for the session because the fee tab re-reads it on every rerun,
    and that is every keystroke in a fee table. The Firm Profile editor
    clears the cache when it saves, so an edited rate card takes effect
    immediately rather than at the next login."""
    if "_firm_rate_card_cache" in st.session_state:
        return st.session_state["_firm_rate_card_cache"]
    card = {}
    try:
        card = firm_profile.rate_card(firm_profile.get_or_create(_fee_history_user_id())) or {}
    except Exception as exc:  # noqa: BLE001 -- an empty rate card is a fine answer
        print(f"[fee prefill] couldn't read the rate card: {exc}", file=sys.stderr)
    st.session_state["_firm_rate_card_cache"] = card
    return card


def _prefill_rates_from_firm_profile() -> int:
    """Fill in each discipline's RATE from the firm's rate card, leaving hours
    alone. Returns how many rows were filled.

    Only ever fills a row whose rate is still zero. A rate the user has typed
    is their pricing decision for this bid, and a standing rate card is not
    entitled to overwrite it."""
    card = {str(k).strip().lower(): float(v or 0) for k, v in _firm_rate_card().items()}
    if not card:
        return 0
    filled = 0
    for line in st.session_state.get("discipline_fee_lines") or []:
        if (line.rate_per_hour or 0) > 0:
            continue
        rate = card.get((line.discipline or "").strip().lower())
        if rate and rate > 0:
            line.rate_per_hour = rate
            filled += 1
    if filled:
        st.session_state._discipline_fee_editor_version += 1
    return filled


def _fee_history_panel(apply_key: str) -> None:
    """The firm's own median split, shown ABOVE the bundled table because it
    is the better source -- and only once there are enough past bids for a
    median to mean anything (see fee_history.MIN_BIDS_FOR_BENCHMARK)."""
    try:
        history = fee_history.fee_history_benchmarks(
            _fee_history_user_id(), st.session_state.get("project_type") or "")
    except Exception as exc:  # noqa: BLE001 -- an extra benchmark is never worth a traceback
        print(f"[fee history] {exc}", file=sys.stderr)
        return
    if not history["disciplines"]:
        return
    st.markdown(f"**Your firm's history** (median of {history['bids']} bids)")
    st.caption(
        "Your own past splits for this project type, from packs you've exported or archived. "
        "This is a record of how your firm has priced, not a market rate and not a "
        "recommendation -- the range shows how much it has actually varied."
    )
    st.dataframe(
        [
            {
                "Discipline": entry["discipline"],
                "Median %": entry["median_pct"],
                "Range": f"{entry['low_pct']:.1f}-{entry['high_pct']:.1f}%",
                "Bids": entry["bids"],
            }
            for entry in history["disciplines"]
        ],
        use_container_width=True, hide_index=True,
    )
    if st.button("Apply your firm's split", key=apply_key, type="primary"):
        discs = [l.discipline for l in (st.session_state.get("discipline_fee_lines") or [])]
        split, source = fee_history.best_available_split(
            _fee_history_user_id(), st.session_state.get("project_type") or "", discs)
        st.session_state.fee_estimates = [
            fee_estimation_engine.DisciplineFeeEstimate(
                discipline=disc, fee_percentage=split.get(disc, 0.0),
                source=f"{fee_history.SOURCE_HISTORY} -- median of {history['bids']} past bids",
                confidence="Your own data",
            )
            for disc in discs
        ]
        st.rerun()


def _target_fee_prefill(state_prefix: str) -> None:
    """"Pre-fill from target fee": split a target across the disciplines by the
    best split available, and derive hours from each row's rate.

    Only ever fills rows that are still untouched. A row the user has priced
    is the whole point of the sheet, and a prefill that overwrote it would be
    the single worst thing this feature could do -- so it asks first, every
    time, and names what it would change.
    """
    lines = st.session_state.get("discipline_fee_lines") or []
    if not lines:
        return
    analysis = st.session_state.get("analysis")
    cap_amount = fee_estimation_engine._parse_currency(
        getattr(analysis, "fee_cap", "") if analysis is not None else "")
    default = float(cap_amount or 0.0)

    with st.expander("Pre-fill from a target fee", expanded=False):
        st.caption(
            "Splits a target across the disciplines below using the best benchmark available "
            "(your firm's history where you have it, otherwise the bundled rule-of-thumb), "
            "then derives hours from each row's rate. It is a first-pass split of a number you "
            "chose -- adjust every row before export."
            + (f" Defaulted to the brief's stated fee cap ({analysis.fee_cap})."
               if cap_amount else "")
        )
        target = st.number_input(
            "Target fee ($, excl. GST)", min_value=0.0, step=1000.0, value=default,
            key=f"{state_prefix}_target_fee",
        )
        priced = [l.discipline for l in lines if l.fee_amount > 0]
        overwrite = False
        if priced:
            st.warning(
                "**Already priced: " + ", ".join(priced) + ".** Pre-filling leaves those rows "
                "exactly as they are unless you tick the box below."
            )
            overwrite = st.checkbox(
                "Also replace the rows I've already priced", value=False,
                key=f"{state_prefix}_target_overwrite",
            )
        if st.button("Pre-fill from target fee", key=f"{state_prefix}_target_apply",
                     type="primary", disabled=target <= 0):
            split, source = fee_history.best_available_split(
                _fee_history_user_id(), st.session_state.get("project_type") or "",
                [l.discipline for l in lines])
            card = {str(k).strip().lower(): float(v or 0) for k, v in _firm_rate_card().items()}
            filled, no_rate = 0, []
            for line in lines:
                if line.fee_amount > 0 and not overwrite:
                    continue
                amount = target * split.get(line.discipline, 0.0) / 100
                if amount <= 0:
                    continue
                rate = line.rate_per_hour or card.get((line.discipline or "").strip().lower(), 0)
                if not rate:
                    # No rate anywhere means hours can't be derived. Leaving
                    # hours at 0 and saying so beats inventing a rate, which
                    # would put a fabricated number in the one table that is
                    # meant to be entirely the user's own.
                    no_rate.append(line.discipline)
                    continue
                line.rate_per_hour = rate
                line.total_hours = round(amount / rate, 1)
                line.note = (line.note or "") or f"Indicative -- {source} split of your target fee"
                filled += 1
            st.session_state._discipline_fee_editor_version += 1
            if filled:
                st.success(
                    f"Pre-filled {filled} discipline(s) from a {source} split. Every figure is "
                    "indicative -- adjust before export."
                )
            if no_rate:
                st.warning(
                    "**No rate for " + ", ".join(no_rate) + "**, so hours couldn't be derived. "
                    "Enter a rate (or add one to your Firm Profile rate card) and pre-fill again."
                )
            if not filled and not no_rate:
                st.info("Nothing to pre-fill -- every row is already priced.")
            st.rerun()


def _fee_ai_refresh(disciplines: list[str], fee_cap: str | None):
    """The AI-modelled benchmark call, fed this brief's real context.

    Returns (estimates, error). Estimates come back EMPTY on failure, and the
    caller must leave the table alone -- the previous version fell back to the
    bundled figures, so a failed call looked exactly like a successful one
    that happened to agree."""
    analysis = st.session_state.get("analysis")
    return fee_estimation_engine.refresh_estimate_from_ai(
        st.session_state.get("project_type") or "",
        disciplines,
        fee_cap,
        st.session_state.ai_config,
        scope_summary=(getattr(analysis, "project_scope", "") or "") if analysis is not None else "",
        analysis=analysis,
    )


def _org_model_from_state():
    """The one org-chart model every output renders from -- preview, pack and
    companion deck -- so an export cannot show a different team from the one
    the user approved on screen."""
    return org_chart_render.build_model(
        st.session_state.get("resource_plan") or [],
        st.session_state.get("client_name") or "",
        st.session_state.get("project_name") or "",
        st.session_state.get("tender_name") or "",
    )


def _org_signature(style: str) -> tuple:
    """Everything the drawn chart depends on. Anything missing here is
    something a change to could leave a stale picture on screen.

    Also reused (see 80_export.py) as the cache signature for the exported
    org chart PPTX, which DOES vary with output_language now (Audit Round 2,
    Part 5) -- included here so a language switch busts that cache instead
    of serving a stale-language deck."""
    plan = st.session_state.get("resource_plan") or []
    return (
        style,
        tuple((a.slot, a.slot_kind, a.person_name, a.is_lead, a.custom_title,
               a.qualification, a.rpeq_status, a.peer_reviewer) for a in plan),
        st.session_state.get("client_name") or "",
        st.session_state.get("project_name") or "",
        st.session_state.get("tender_name") or "",
        st.session_state.get("proposal_theme") or "",
        st.session_state.get("output_language") or "en",
    )


_ORG_PNG_CACHE_SIZE = 8


def _org_png(style: str | None = None) -> bytes | None:
    """The current org chart drawn in `style`, cached on its content
    signature. Falls back to the original PIL renderer (modules/org_chart.py)
    if the styled renderer fails -- a chart the reader can follow beats no
    chart at all. None only when both fail; every caller handles that."""
    style = style or st.session_state.get("org_chart_style") or org_chart_render.DEFAULT_STYLE
    signature = _org_signature(style)
    cache = st.session_state.setdefault("_org_png_cache", {})
    if signature in cache:
        return cache[signature]
    png = None
    try:
        accent = f"#{export_docx._theme_colours(st.session_state.get('proposal_theme'))['accent']}"
        png = org_chart_render.render_png(_org_model_from_state(), style, accent)
    except Exception as exc:  # noqa: BLE001 -- a preview is never worth a traceback
        print(f"[org chart preview] {exc}", file=sys.stderr)
    if png is None:
        try:
            png = org_chart.render_org_chart(
                st.session_state.get("resource_plan") or [],
                theme_name=st.session_state.get("proposal_theme"),
                project_title=(st.session_state.get("tender_name")
                               or st.session_state.get("project_name") or None),
            )
        except Exception as exc:  # noqa: BLE001
            print(f"[org chart fallback] {exc}", file=sys.stderr)
    if len(cache) >= _ORG_PNG_CACHE_SIZE:
        cache.pop(next(iter(cache)))
    cache[signature] = png
    return png


def _org_chart_style_control() -> None:
    """The four-way org-chart style picker, with a live preview drawn from
    this project's own team.

    A preview rather than four stock thumbnails, for the same reason the
    program styles get one: the styles differ most in how they cope with THIS
    project's role names and discipline count, which a generic sample cannot
    show."""
    st.markdown("**Org chart style**")
    st.caption(
        "How the chart is drawn in the exported pack and the companion PowerPoint. "
        "Purely presentational -- it never changes who is on the team."
    )
    st.radio(
        "Org chart style", org_chart_render.STYLES,
        key="org_chart_style", horizontal=True, label_visibility="collapsed",
        format_func=lambda s: org_chart_render.STYLE_LABELS.get(s, s),
    )
    st.caption(org_chart_render.STYLE_DESCRIPTIONS.get(st.session_state.org_chart_style, ""))


def _methodology_signature() -> tuple:
    """Everything the drawn methodology preview depends on -- the reviewed
    stage grid (or, before that exists, the brief's scope items via
    `analysis`), the program's week labels (the date chevrons), and the
    client/project names shown in the title. Anything missing here is
    something a change to could leave a stale preview on screen."""
    stages = st.session_state.get("methodology_stages") or []
    return (
        tuple((s.name, tuple(s.key_tasks), tuple(s.engagement_activities), s.outcome,
              tuple(s.deliverables), s.week_start, s.week_end) for s in stages),
        tuple(st.session_state.get("program_week_labels") or []),
        st.session_state.get("client_name") or "",
        st.session_state.get("project_name") or "",
    )


_METHODOLOGY_PNG_CACHE_SIZE = 8


def _methodology_png(style: str | None = None) -> bytes | None:
    """The current methodology, drawn in `style`, cached on its content
    signature -- same pattern as _org_png/_program_png. None only if the
    renderer itself fails; every caller handles that."""
    style = style or st.session_state.get("methodology_style") or methodology_render.DEFAULT_STYLE
    signature = (style,) + _methodology_signature()
    cache = st.session_state.setdefault("_methodology_png_cache", {})
    if signature in cache:
        return cache[signature]
    png = None
    try:
        png = methodology_render.render_png(
            st.session_state.get("analysis"), style=style,
            stages=st.session_state.get("methodology_stages"),
            week_labels=st.session_state.get("program_week_labels"),
            client_name=st.session_state.get("client_name") or "",
            project_name=st.session_state.get("project_name") or "",
        )
    except Exception as exc:  # noqa: BLE001 -- a preview is never worth a traceback
        print(f"[methodology preview] {exc}", file=sys.stderr)
    if len(cache) >= _METHODOLOGY_PNG_CACHE_SIZE:
        cache.pop(next(iter(cache)))
    cache[signature] = png
    return png


def _methodology_style_control() -> None:
    """The four-way methodology presentation style picker, with a live
    preview drawn from this project's own reviewed stage grid (or the
    pre-stages boilerplate when the stage drafter hasn't been run yet)."""
    st.markdown("**Methodology presentation style**")
    st.caption(
        "How the methodology table is drawn in the companion PowerPoint. Purely "
        "presentational -- it never changes the tasks, activities, outcomes or "
        "deliverables themselves."
    )
    st.radio(
        "Methodology presentation style", methodology_render.STYLES,
        key="methodology_style", horizontal=True, label_visibility="collapsed",
        format_func=lambda s: methodology_render.STYLE_LABELS.get(s, s),
    )
    st.caption(methodology_render.STYLE_DESCRIPTIONS.get(st.session_state.methodology_style, ""))
    _preview = _methodology_png()
    if _preview:
        st.image(_preview, use_container_width=True)
    else:
        st.caption(
            "Couldn't draw the preview just now -- your methodology content is unaffected."
        )


def _management_insert_index(plan: list, role: str) -> int:
    """Where a re-added management role goes back in the plan.

    Before the first management row that comes AFTER it in
    resourcing.MANDATORY_ORG_ROLES, so the chain reads client -> director ->
    manager -> design manager however many times it has been removed and
    restored. Appending instead would leave the restored role below the
    discipline leads in the pen-pic order."""
    order = resourcing.MANDATORY_ORG_ROLES
    if role not in order:
        return len(plan)
    rank = order.index(role)
    for index, assignment in enumerate(plan):
        if getattr(assignment, "slot_kind", "") != "management":
            continue
        slot = getattr(assignment, "slot", "")
        if slot in order and order.index(slot) > rank:
            return index
    # No later management row -- go after the last management row there is.
    last = [i for i, a in enumerate(plan) if getattr(a, "slot_kind", "") == "management"]
    return (last[-1] + 1) if last else 0


def _render_resource_rows(kind: str, known_names: list) -> None:
    """Render the assign-a-person rows for the Team & Resourcing tab, for either
    the 'management' roles or the 'discipline' leads. Each row lets the user pick
    a known (CV-derived or typed) name or choose '(type a name)' to enter someone
    ad hoc; discipline rows can also be removed. Mutates st.session_state.resource_plan
    in place.

    For 'discipline' rows only, a lead can also carry any number of support
    members underneath it (e.g. "Ryan Swagemakers" added under the "Structural"
    lead, "Mat Williams") -- each with their own free-text title for this
    project (e.g. "Bridge Engineer"), since a support member's title is rarely
    just the discipline name. Support rows always render immediately after
    their lead (see resourcing.normalize_plan_disciplines, which keeps the plan
    grouped that way), indented with an "under <lead>" caption."""
    TYPE_SENTINEL = "— type a name —"
    UNASSIGNED = "(unassigned)"
    options = [UNASSIGNED] + known_names + [TYPE_SENTINEL]
    remove_index = None
    add_support_after = None  # index of the lead row to add a new support row under
    plan = st.session_state.resource_plan
    for i, a in enumerate(plan):
        if a.slot_kind != kind:
            continue
        is_support = kind == "discipline" and not a.is_lead
        # A management row gets a remove control only if the role is optional
        # (resourcing.OPTIONAL_MANAGEMENT_ROLES) -- a chart with no project
        # director isn't a chart.
        removable_management = kind == "management" and resourcing.is_removable_management_role(a.slot)
        cols = (st.columns([3, 3, 1, 1]) if kind == "discipline"
                else (st.columns([3, 3, 1]) if removable_management else st.columns([3, 3])))
        with cols[0]:
            if is_support:
                st.caption(f"↳ under {a.slot}")
                a.custom_title = st.text_input(
                    "Title", value=a.custom_title, key=f"res_title_{kind}_{i}",
                    label_visibility="collapsed",
                    placeholder="Their title on this project, e.g. Bridge Engineer",
                )
            else:
                st.markdown(f"**{a.slot}**")
        with cols[1]:
            current = (a.person_name or "").strip()
            if current and current in known_names:
                idx = options.index(current)
            elif current:
                idx = options.index(TYPE_SENTINEL)
            else:
                idx = 0
            choice = st.selectbox(
                "Assigned to", options, index=idx,
                key=f"res_sel_{kind}_{i}", label_visibility="collapsed",
            )
            if choice == TYPE_SENTINEL:
                typed = st.text_input(
                    "Name", value=current if current not in known_names else "",
                    key=f"res_txt_{kind}_{i}", label_visibility="collapsed",
                    placeholder="Type the person's name",
                )
                a.person_name = typed.strip()
                a.from_cv = False
            elif choice == UNASSIGNED:
                a.person_name = ""
                a.from_cv = False
            else:
                a.person_name = choice
                a.from_cv = choice in known_names
        if removable_management:
            with cols[2]:
                if st.button(
                    "✕", key=f"res_del_{kind}_{i}",
                    help=f"Remove {a.slot} from this project -- not every commission has one. "
                         f"You can add it back below.",
                 type="primary"):
                    remove_index = i
        if kind == "discipline":
            with cols[2]:
                if not is_support and st.button(
                    "+ member", key=f"res_addsup_{kind}_{i}", help="Add a team member under this lead",
                 type="primary"):
                    add_support_after = i
            with cols[3]:
                if st.button(
                    "✕", key=f"res_del_{kind}_{i}",
                    help="Remove this discipline (and anyone added under it)" if not is_support else "Remove this team member",
                 type="primary"):
                    remove_index = i
            if not is_support:
                # One optional peer-reviewer name per discipline -- the org
                # chart's Peer Review element (all four styles) reads this
                # straight off the lead row. Blank renders as a red TBC on
                # the chart, never a guess or a silently-omitted row.
                pr_col1, pr_col2 = st.columns([1, 4])
                with pr_col1:
                    st.caption(f"↳ {a.slot} peer reviewer")
                with pr_col2:
                    a.peer_reviewer = st.text_input(
                        f"Peer reviewer for {a.slot} (optional)", value=a.peer_reviewer,
                        key=f"res_peerrev_{kind}_{i}", label_visibility="collapsed",
                        placeholder="Optional -- who peer-reviews this discipline's work. Blank shows a red TBC on the chart.",
                    )
    if add_support_after is not None:
        lead = plan[add_support_after]
        # Insert right after the lead's existing block (the lead plus any
        # support rows already under it), so the plan stays grouped/contiguous
        # for every downstream reader (org chart, Small Scope Project Team,
        # resourcing.discipline_groups/normalize_plan_disciplines).
        insert_at = add_support_after + 1
        while (insert_at < len(plan) and plan[insert_at].slot_kind == "discipline"
               and not plan[insert_at].is_lead and plan[insert_at].slot == lead.slot):
            insert_at += 1
        plan.insert(insert_at, resourcing.ResourceAssignment(slot=lead.slot, slot_kind="discipline", is_lead=False))
        st.rerun()
    if remove_index is not None:
        removed = plan.pop(remove_index)
        if removed.slot_kind == "management":
            # Recorded, not merely absent -- otherwise the next reconcile pass
            # (or a re-run of Tender Analysis, which rebuilds the plan) puts it
            # straight back. Same reasoning as dismissed_disciplines.
            if removed.slot not in st.session_state.removed_management_roles:
                st.session_state.removed_management_roles.append(removed.slot)
        elif removed.is_lead:
            # Cascade: a removed lead's support rows have nothing left to be
            # nested under, so they go too rather than becoming orphans.
            plan[:] = [x for x in plan if not (x.slot_kind == "discipline" and not x.is_lead and x.slot == removed.slot)]
            # Remember the removal so the brief re-sync doesn't immediately re-add it.
            label = resourcing.canonical_discipline(removed.slot)
            if label and label not in st.session_state.dismissed_disciplines:
                st.session_state.dismissed_disciplines.append(label)
        st.rerun()


# The six "Done entering data -- refresh ..." deferred-apply fee tables (see
# _render_large_discipline_fee_table() and its siblings) each track their own
# "*_last_applied_editor_sig" -- None means "never applied yet, bypass the
# tick requirement on the very next render." That bypass can fire once
# *before* a project is ever loaded (e.g. against the empty/default project
# these tables briefly render against right after account creation, seeded
# with nothing but the always-included "Project Management" line), which
# permanently consumes it with a near-empty baseline. Loading a real project
# afterwards then makes every one of these tables look permanently "pending"
# (mismatched against that stale one-row baseline) even though the user
# hasn't touched anything -- confirmed via a debug probe showing exactly this
# for the large-scope "Indicative fee split by discipline" table. Resetting
# all of them here, alongside the loaded values, gives a freshly (re)loaded
# project a genuine first-load bypass keyed to its own real data.
_FEE_TABLE_APPLY_STATE_PREFIXES = (
    "_disc_fee_", "_letter_disc_fee_",
    "_scope_fee_", "_large_scope_fee_",
    "_pct_fee_", "_letter_pct_fee_",
)


def _apply_loaded_project(loaded_state: dict, source_label: str) -> None:
    """Shared by both the local 'Open' button and the manual zip uploader --
    overwrites every project_store-managed session_state key with the loaded
    values, resets save/export bookkeeping that no longer applies to the
    newly-loaded project, and reruns so every tab reflects it immediately."""
    for k, v in loaded_state.items():
        st.session_state[k] = v
    # Round 3, Part 3: _last_draft_metered_signature now travels with the
    # project (see project_store.PLAIN_KEYS), so a normal save/load round
    # trip preserves the metering baseline and a page refresh can no longer
    # reset a paid project back to "first free generation". But an OLDER
    # project file saved before this key existed -- or any other path that
    # lands here with drafts already generated but no baseline in the
    # payload -- must NOT be silently treated as a free first-ever
    # generation (that's exactly what a missing baseline normally signals,
    # see _draft_would_consume_pass()). Stamp the CURRENT signature as the
    # baseline right here instead, so the very next regenerate compares
    # against what was actually loaded, not against nothing.
    if st.session_state.get("drafts") and not st.session_state.get("_last_draft_metered_signature"):
        st.session_state["_last_draft_metered_signature"] = _draft_generation_input_signature()
    st.session_state._project_save_bytes = None
    st.session_state.docx_buffer = None
    for prefix in _FEE_TABLE_APPLY_STATE_PREFIXES:
        st.session_state[f"{prefix}last_applied_editor_sig"] = None
    # The "re-enter your AI provider settings" follow-up only applies in the
    # desktop/BYOK build, where ai_config lives purely in session_state and
    # loading a project doesn't touch it, but a fresh session might not have
    # a key typed in yet. In SaaS mode ai_config is auto-filled from the
    # server-side ANTHROPIC_API_KEY at session start (see _ENV_ANTHROPIC_KEY)
    # and never needs re-entering, so telling a SaaS customer to go do that
    # in a sidebar field that doesn't exist for them was just wrong.
    if IS_SAAS_MODE:
        st.success(i18n.t("chrome_loaded_project_success", source=source_label))
    else:
        st.success(i18n.t("chrome_loaded_project_reenter_ai_success", source=source_label))
    st.rerun()


def _maybe_autosave() -> None:
    """Called once per script run, after every tab has had a chance to mutate
    session_state. Debounced by AUTOSAVE_INTERVAL_SECONDS rather than saving
    on literally every rerun (Streamlit reruns on most widget interactions,
    including every single data_editor cell edit in the Fees & Program tab --
    writing a multi-MB zip to disk that often would be wasteful and could lag
    the UI). Silently skipped if there's no project name yet, so a blank
    session doesn't create a stray 'untitled_project' file.

    In SAAS_MODE, saves to the database under the logged-in user's account
    (cloud_project_store) instead of the server's local disk -- otherwise
    every user's in-progress work (uploaded briefs, drafts, team CVs) exists
    ONLY in that one browser tab's live session, and is lost outright on any
    page refresh, dropped connection, or redeploy, with no way to recover it
    (see the "My projects" sidebar section, which reads from the same
    table). local_project_store is kept for the non-SaaS local prototype."""
    if not st.session_state._autosave_enabled:
        return
    project_id = _project_identifier()
    if not project_id:
        # No project/tender name yet -- this used to skip autosave entirely
        # until one was entered, which meant a user who'd already uploaded a
        # real tender brief (the slow, easy-to-forget step) but simply
        # hadn't gotten to Project Setup yet would lose that brief outright
        # on any refresh, dropped connection, or accidental tab close --
        # exactly the state a rushing trial user is in. Once there's real
        # content worth not losing, fall back to a fixed "Untitled project"
        # slot instead of skipping -- multiple different unnamed projects in
        # the same account will share and overwrite that one slot until
        # each is given its own name, but that's an acceptable trade-off
        # for a "don't lose today's work" safety net, not a version
        # history. Still skipped entirely for a genuinely blank session (no
        # name AND no uploaded brief yet) so opening the app doesn't create
        # a stray entry before there's anything worth saving.
        if st.session_state.tender_extracted is None:
            return
        project_id = "Untitled project"
    now = time.time()
    if now - st.session_state._last_autosave_ts < AUTOSAVE_INTERVAL_SECONDS:
        return
    try:
        # Serialize once up front (this is the same work save_cloud()/
        # save_local() used to each do internally) so its content hash can
        # be compared against the last save's hash -- if nothing that
        # save_project() actually captures has changed since then, skip the
        # write entirely. Hashing costs something every interval too, but
        # it's in-process and cheap; the write it can now skip is the
        # expensive part at scale (for cloud_project_store specifically: a
        # network round trip writing a multi-MB blob to Postgres, on every
        # interval, for every active user, whether or not anything in it
        # actually changed).
        blob = project_store.save_project(st.session_state)
        content_hash = hashlib.sha256(blob).hexdigest()
        if content_hash == st.session_state.get("_last_autosave_hash"):
            st.session_state._last_autosave_ts = now
            return
        if IS_SAAS_MODE and current_user:
            slug = cloud_project_store.save_cloud(current_user.id, project_id, blob)
            st.session_state._last_autosave_ts = now
            st.session_state._last_autosave_path = slug
            st.session_state._last_autosave_error = ""
        else:
            path = local_project_store.save_local(project_id, blob)
            st.session_state._last_autosave_ts = now
            st.session_state._last_autosave_path = path
            st.session_state._last_autosave_error = ""
        st.session_state["_last_autosave_hash"] = content_hash
    except Exception as exc:
        # Auto-save is a convenience, not a step the user is waiting on -- a failure here
        # (e.g. disk full, folder permissions, a transient DB hiccup) shouldn't interrupt
        # whatever they were doing, so this still doesn't raise/stop the script. But it
        # used to swallow the error completely (bare `except: pass`), which is exactly
        # how a real problem -- like the DATABASE_URL misconfiguration found and fixed
        # earlier -- could silently fail every single save with the user having no way
        # to know, only discovering it later when a project they thought was saved
        # wasn't there. Recording it here lets the "My projects" section surface a
        # visible warning instead (see the sidebar code that reads
        # _last_autosave_error), and _last_autosave_ts is deliberately NOT updated on
        # failure, so the very next rerun retries immediately rather than waiting out
        # the normal debounce interval.
        st.session_state._last_autosave_error = str(exc) or exc.__class__.__name__


def _ensure_divider_config(sections) -> None:
    """Give every current section a sensible default divider design (layout + which
    uploaded photo, if any) the first time it's seen, without clobbering choices the
    user already made. Called each time the Graphics & Design tab renders."""
    photos = st.session_state.project_photo_bytes
    config = st.session_state.section_divider_config
    for i, s in enumerate(sections):
        if s.title not in config:
            layout = "Photo + gradient" if photos else "Solid colour"
            config[s.title] = {
                "layout": layout,
                "photo_index": (i % len(photos)) if photos else None,
                "quote_index": None,
                "photo_caption": "",
            }
        else:
            # Back-fill for configs created before "photo_caption" existed
            # (older saved/autosaved projects) -- never clobber a value
            # that's already there.
            config[s.title].setdefault("photo_caption", "")
    # Drop config for sections that no longer exist (e.g. after a structure rebuild).
    current_titles = {s.title for s in sections}
    for stale in [t for t in config if t not in current_titles]:
        del config[stale]


_init_state()

# Seed a NEW project's blank fields from the firm profile -- bidder name,
# signatory block, standing terms. Once per session, not per rerun: seeding
# on every rerun would silently refill a field the user had deliberately
# cleared for this bid, which is exactly the kind of quiet wrongness a
# proposal tool must not have. Loading a saved project overwrites these
# fields from the file anyway, so the two paths don't fight.
_ensure_photo_ids()

if not st.session_state.get("_firm_profile_seeded"):
    try:
        _seed_project_from_firm_profile()
    except Exception:
        pass
    st.session_state._firm_profile_seeded = True

# Attribute every AI call made inline in THIS script run (i.e. not via the
# job queue -- that path carries its own usage_context, see
# _run_job_or_inline) to the logged-in user and current project, for
# per-bid cost logging (db.AiCallLog). Re-set on every rerun because the
# project fields can change between runs. Best-effort on purpose.
try:
    ai_interface.set_usage_context(
        user_id=current_user.id if (IS_SAAS_MODE and current_user) else None,
        project_key=_current_project_key(),
        project_name=st.session_state.get("project_name", ""),
    )
except Exception:
    pass


